# export_excel.py — ВЫГРУЗКА в Excel
# ---------------------------------------------------------------
# Собирает упрощённую книгу QA_System.xlsx из текущих данных.
# ---------------------------------------------------------------
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from config import COMPS, GRADES, PPL, HISTORY, SCORES, GOALS, BASES
from grading import build_results

H_FONT = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="305496")
SIDE = Side(style="thin", color="BFBFBF")
BD = Border(left=SIDE, right=SIDE, top=SIDE, bottom=SIDE)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _head(ws, r, headers):
    for i, t in enumerate(headers, 1):
        c = ws.cell(r, i, t)
        c.font = H_FONT
        c.fill = FILL
        c.alignment = CENTER
        c.border = BD


def _row(ws, r, values):
    for i, v in enumerate(values, 1):
        c = ws.cell(r, i, v)
        c.border = BD
        c.alignment = CENTER


def export(path="QA_System.xlsx"):
    wb = Workbook()

    # Summary / Дашборд
    ws = wb.active
    ws.title = "Дашборд"
    results = build_results()
    headers = ["Сотрудник", "% грейд", "Грейд", "% премии", "Оценка",
               "% цели", "Премия", "Коэф.", "Статус", "Оклад"]
    _head(ws, 1, headers)
    for i, rc in enumerate(results):
        _row(ws, i + 2, [rc[h] for h in headers])
        ws.cell(i + 2, 4).number_format = "0%"
    n = len(results) + 3
    ws.cell(n, 1, "Фонд премий, руб.:").font = Font(bold=True)
    ws.cell(n, 2, sum(r["Премия"] for r in results))

    # Компетенции
    ws2 = wb.create_sheet("Компетенции")
    _head(ws2, 1, ["Компетенция", "Категория", "Приоритет", "Вес", "В грейд?"])
    for i, c in enumerate(COMPS):
        _row(ws2, i + 2, list(c))

    # Оценки
    ws3 = wb.create_sheet("Оценки")
    _head(ws3, 1, ["Компетенция"] + PPL)
    for idx, c in enumerate(COMPS):
        _row(ws3, idx + 2, [c[0]] + [SCORES[p][idx] for p in PPL])

    # Грейды
    ws4 = wb.create_sheet("Грейды")
    _head(ws4, 1, ["Грейд", "Порог %", "% премии"])
    for i, (g, thr, pc) in enumerate(GRADES):
        _row(ws4, i + 2, [g, thr, pc])
        ws4.cell(i + 2, 3).number_format = "0%"

    wb.save(path)
    return path
